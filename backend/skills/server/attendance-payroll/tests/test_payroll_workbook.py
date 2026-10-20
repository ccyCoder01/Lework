from __future__ import annotations

import copy
import calendar
import hashlib
import json
import sys
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
import build_payroll_workbook as exporter  # noqa: E402
import project_binding  # noqa: E402


COMPONENTS = exporter.COMPONENTS
MONTH = "2026-06"
BATCH = "task_A12-2026-06"


def _fixture(count: int = 31, month: str = MONTH) -> dict:
    year, month_number = map(int, month.split("-"))
    days_in_month = calendar.monthrange(year, month_number)[1]
    batch = f"task_A12-{month}"
    snapshot_records = []
    calc_items = []
    review_people = []
    for i in range(1, count + 1):
        name = f"合成人员{i:02d}"
        record_id = f"record-{i:02d}"
        review_people.append({"record_id": record_id, "name": name})
        snapshot_records.append({
            "record_id": record_id, "payroll_person_id": i, "roster_id": i + 100, "name_snapshot": name,
            "employee_type_snapshot": 1, "company_snapshot": "合成公司",
            "project_snapshot": "合成项目", "source_project_snapshot": "合成项目",
            "position_snapshot": "合成岗位", "source_actual_work_days": 20,
            "derived_actual_work_days": 20, "confirmed_actual_work_days": 20,
            "personal_leave_days": 0, "sick_leave_days": 0, "absent_days": 0,
            "reconciliation_status": 1, "remark": None, "items": [
                {"day": day, "status": 1 if day <= min(20, days_in_month) else 5, "value": None,
                 "raw_value": "8" if day <= 20 else "休"}
                for day in range(1, days_in_month + 1)
            ],
        })
        components = [{"code": code, "name": label, "source_present": True,
                       "source_amount": "0.00" if index == 0 else "11.25",
                       "calculated_amount": "0.00" if index == 0 else "11.25"}
                      for index, (code, label) in enumerate(COMPONENTS)]
        calc_items.append({
            "batch_id": batch, "payroll_person_id": i, "roster_id": i + 100,
            "name": name, "employee_type": 1, "project_name": "合成项目", "company_name": "合成公司", "position": "合成岗位",
            "attendance_source": "BATCH_SNAPSHOT", "attendance_batch_operation_id": f"submit:{batch}",
            "his_salary_record_id": i + 200, "historical_source_month": month, "components": components,
            "historical_gross_amount": "90.00", "calculated_gross_amount": "90.00", "difference": "0.00",
        })
    source_people = []
    for i in range(1, count + 1):
        source_people.append({
            "record_id": f"record-{i:02d}", "name": f"合成人员{i:02d}", "project": "合成项目",
            "company_name": "合成公司", "employee_type": 1, "position": "合成岗位", "printed_work_days": 20,
            "row_binding_status": "VERIFIED", "row_index": i, "printed_work_days_row_index": i,
            "daily": [{"day": day, "status": "work" if day <= min(20, days_in_month) else "rest",
                       "raw_mark": "8" if day <= 20 else "休"} for day in range(1, days_in_month + 1)],
        })
    return {
        "workflow_state": {"phase": "CALCULATED_ARTIFACT_PENDING", "attendance_generation": 2,
                           "attendance_approved_generation": 2, "human_precheck_confirmed": True,
                           "precheck_confirmed_snapshot_id": "precheck-x", "precheck_snapshot_id": "precheck-x",
                           "calculation_completed": True},
        "latest_precheck": {"snapshot_id": "precheck-x", "valid": True, "batch_blocked": False,
                            "participant_count": count, "calculated_count": count, "blocked_count": 0},
        "attendance_snapshot": {"attendance_month": month, "records": snapshot_records},
        "_source_people": source_people,
        "_source_matches": [
            {"record_id": f"record-{i:02d}", "status": "matched", "payroll_person_id": i, "roster_id": i + 100}
            for i in range(1, count + 1)
        ],
        "calculation_result": {"payroll_month": month, "batch_id": batch, "operation_id": f"calculate:{batch}",
                               "precheck_snapshot_id": "precheck-x", "valid": True, "batch_blocked": False,
                               "calculation_stage": "COMPLETED", "population_count": count,
                               "participant_count": count, "calculated_count": count, "blocked_count": 0,
                               "status": "SUCCESS", "issues": [], "items": calc_items,
                               "summary": {"calculated_count": count, "population_count": count,
                                           "calculated_gross_total": str(90 * count),
                                           "historical_gross_total": str(90 * count), "difference_total": "0.00"}},
    }


def _write_inputs(tmp_path: Path, data: dict) -> tuple[Path, Path, Path]:
    month = data["attendance_snapshot"]["attendance_month"]
    batch = data["calculation_result"]["batch_id"]
    input_path = tmp_path / "payroll-workbook-input.json"
    persisted_data = {key: value for key, value in data.items() if not key.startswith("_")}
    input_path.write_text(json.dumps(persisted_data, ensure_ascii=False), encoding="utf-8")
    attendance_path = tmp_path / "attendance.json"
    attendance_path.write_text(json.dumps({"schema_version": 1, "month": month, "people": data["_source_people"],
                                           "source_warnings": []}, ensure_ascii=False), encoding="utf-8")
    digest = hashlib.sha256(attendance_path.read_bytes()).hexdigest()
    roster_path, binding_path = project_binding.default_paths(attendance_path)
    roster_path.parent.mkdir(parents=True, exist_ok=True)
    snapshot_by_record = {row["record_id"]: row for row in data["attendance_snapshot"]["records"]}
    source_to_target = {
        person["project"]: snapshot_by_record.get(person["record_id"], {}).get("project_snapshot", person["project"])
        for person in data["_source_people"]
    }
    target_projects = sorted(set(source_to_target.values()))
    roster_path.write_text(json.dumps({"count": len(target_projects), "items": [
        {"roster_id": 100 + i, "name": f"合成人员{i + 1:02d}", "project_name": project,
         "company_name": "合成公司", "employee_type": 1, "position": "合成岗位"}
        for i, project in enumerate(target_projects)
    ]}, ensure_ascii=False), encoding="utf-8")
    prepared_binding = project_binding.prepare(attendance_path, roster_path)
    binding_request = tmp_path / "project-binding-request.json"
    bindings = [{"source_project_name": source, "payroll_project_name": target}
                for source, target in sorted(source_to_target.items())]
    display = (f"考勤月份：{month}\n考勤人数：{len(data['_source_people'])} 人\n项目对应关系：\n" + "\n".join(
        f"来源项目：{row['source_project_name']} → Payroll 项目：{row['payroll_project_name']}"
        for row in bindings
    ) + "\n请确认考勤事实及以上项目对应关系。")
    confirmation = "确认项目 " + "；".join(
        f"{row['source_project_name']} 对应 {row['payroll_project_name']}" for row in bindings
    )
    digest_json = lambda value: hashlib.sha256(json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"),
    ).encode("utf-8")).hexdigest()
    binding_request.write_text(json.dumps({
        "month": month, "attendance_sha256": digest,
        "candidate_response_sha256": prepared_binding["candidate_response_sha256"], "bindings": bindings,
        "display_evidence": {
            "message_index": 0, "message": display,
            "message_sha256": hashlib.sha256(display.encode("utf-8")).hexdigest(),
            "month": month, "attendance_sha256": digest,
            "candidate_response_sha256": prepared_binding["candidate_response_sha256"],
            "bindings": bindings, "mapping_sha256": digest_json(bindings),
        },
        "user_confirmation": {"message_index": 1, "message": confirmation,
                              "message_sha256": hashlib.sha256(confirmation.encode("utf-8")).hexdigest()},
        "conversation_messages": [
            {"index": 0, "role": "assistant", "content": display},
            {"index": 1, "role": "user", "content": confirmation},
        ],
    }, ensure_ascii=False), encoding="utf-8")
    project_binding.confirm(attendance_path, roster_path, binding_request, binding_path)
    calc_items = persisted_data["calculation_result"]["items"]
    person_ids = sorted(item["payroll_person_id"] for item in calc_items)
    record_ids = [f"record-{pid:02d}" for pid in person_ids]
    calc_names = {item["payroll_person_id"]: item["name"] for item in persisted_data["calculation_result"]["items"]}
    people = [{"record_id": record_id, "payroll_person_id": int(record_id.split('-')[1]),
               "roster_id": int(record_id.split('-')[1]) + 100,
               "name": calc_names[int(record_id.split('-')[1])]}
              for record_id in record_ids]
    scope = {"month": month, "version": "review-v2", "attendance_sha256": digest, "record_ids": record_ids}
    gate_request = {
        "action": "PAYROLL_EXPORT", "user_messages": ["工资核算文件交付"],
        "expected_submit_operation_id": f"submit:{batch}",
        "expected_calculate_operation_id": f"calculate:{batch}",
        "evidence": {
            "attendance_path": str(attendance_path),
            "attendance_review": {"status": "DECLARED", **scope, "people": people,
                                  "visible_after_user_message_index": 0},
            "identity_matching": {"status": "COMPLETE", **scope, "people": people},
            "salary_readiness": {"status": "COVERAGE_CHECKED", **scope, "people": people},
            "submission_scope_review": {"status": "CONFIRMED", **scope, "people": people,
                                        "confirmation_message_index": 0},
            "batch_context": {"status": "VERIFIED_NO_CONFLICT", "task_id": "task_A12", "batch_id": batch,
                              "operation_id": f"submit:{batch}", **scope,
                              "payroll_person_ids": person_ids},
            "submission_receipt": {"status": "SUCCEEDED", **scope, "batch_id": batch,
                                   "operation_id": f"submit:{batch}", "payload_sha256": "a" * 64,
                                   "payroll_person_ids": person_ids},
            "payroll_precheck": {"status": "VALID", "valid": True, **scope, "people": people,
                                 "batch_id": batch, "operation_id": f"precheck:{batch}"},
            "calculation_receipt": {"status": "SUCCEEDED", **scope, "people": people, "batch_id": batch,
                                    "operation_id": f"calculate:{batch}"},
            "payroll_operation_state": {
                "status": "VERIFIED", "task_id": "task_A12", "batch_id": batch, "month": month,
                "operations": {
                    "submission": {"operation_id": f"submit:{batch}", "status": "SUCCEEDED",
                                   "month": month, "attendance_sha256": digest, "record_ids": record_ids,
                                   "payroll_person_ids": person_ids},
                    "calculation": {"operation_id": f"calculate:{batch}", "status": "SUCCEEDED",
                                    "month": month, "attendance_sha256": digest, "record_ids": record_ids,
                                    "payroll_person_ids": person_ids},
                },
            },
        },
    }
    gate_path = tmp_path / "gate.json"
    gate_path.write_text(json.dumps(gate_request, ensure_ascii=False), encoding="utf-8")
    calculation_response_path = tmp_path / "calculation-response.json"
    calculation_response_path.write_text(json.dumps(persisted_data["calculation_result"], ensure_ascii=False), encoding="utf-8")
    matches_path = tmp_path / "person-matches.json"
    matches_path.write_text(json.dumps({"schema_version": 1,
                                        "matching_rules_version": exporter.person_matches.MATCHING_RULES_VERSION,
                                        "attendance_sha256": digest,
                                        "project_binding_sha256": json.loads(binding_path.read_text(encoding="utf-8"))["binding_sha256"],
                                        "matches": data["_source_matches"]},
                                       ensure_ascii=False), encoding="utf-8")
    selected = set(person_ids)
    submission_records = [
        {key: record.get(key) for key in exporter.SUBMISSION_RECORD_KEYS}
        for record in persisted_data["attendance_snapshot"]["records"]
        if record.get("payroll_person_id") in selected
    ]
    submission_path = tmp_path / "submission-request.json"
    submission_path.write_text(json.dumps({"attendance_month": month, "submission_mode": "BATCH_SNAPSHOT",
                                           "batch_id": batch, "operation_id": f"submit:{batch}",
                                           "records": submission_records}, ensure_ascii=False), encoding="utf-8")
    return input_path, gate_path, attendance_path


def _call_export(input_path: Path, gate_path: Path, output: Path) -> dict:
    root = input_path.parent
    return exporter.export(input_path, gate_path, output, root / "calculation-response.json",
                           root / "person-matches.json", root / "submission-request.json")


def _read_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def _write_json(path: Path, value: dict, **kwargs) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, **kwargs), encoding="utf-8")


def _export(tmp_path: Path, data: dict | None = None, name: str = "out.xlsx") -> Path:
    input_path, gate_path, _ = _write_inputs(tmp_path, data or _fixture())
    output = tmp_path / name
    result = _call_export(input_path, gate_path, output)
    assert result["ok"] is True
    return output


def test_success_has_fixed_five_sheets_and_historical_headers(tmp_path: Path):
    path = _export(tmp_path)
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == list(exporter.SHEETS)
    assert [wb[name].max_column for name in wb.sheetnames] == [19, 25, 44, 8, 4]
    assert [list(next(wb[name].iter_rows(min_row=3, max_row=3, values_only=True))) for name in wb.sheetnames] == [
        ["姓名", "项目", "公司/部门", "岗位", "人员类别", "基础工资", "绩效工资", "工龄工资", "职称工资", "施工补贴",
         "话费补贴", "降温费", "交通补贴", "双休日加班工资", "历史工资来源月份", "历史应发工资", "核算应发工资", "与历史差额", "状态"],
        ["姓名", "项目", "公司/部门", "岗位", "人员类别", "历史工资来源月份", "历史应发工资", "基础工资", "基础工资来源状态",
         "绩效工资", "绩效工资来源状态", "工龄工资", "工龄工资来源状态", "职称工资", "职称工资来源状态", "施工补贴",
         "施工补贴来源状态", "话费补贴", "话费补贴来源状态", "降温费", "降温费来源状态", "交通补贴", "交通补贴来源状态",
         "双休日加班工资", "双休日加班工资来源状态"],
        ["姓名", "考勤记录ID", "项目", "公司/部门", "岗位", "人员类别", "来源出勤天数", "派生出勤天数", "确认出勤天数", "事假天数",
         "病假天数", "旷工天数", "对账状态", "备注", *[f"{day}日" for day in range(1, 31)]],
        ["姓名", "项目", "公司/部门", "岗位", "历史工资来源月份", "历史应发工资", "本次核算应发", "与历史差额"],
        ["姓名", "处理状态", "问题类型", "问题说明"],
    ]
    assert [wb[name].freeze_panes for name in wb.sheetnames] == ["A4", "A4", "O4", "A4", "A4"]
    assert all(wb[name].auto_filter.ref.startswith("A3:") for name in wb.sheetnames)
    assert not any(c.data_type == "f" for ws in wb.worksheets for row in ws.iter_rows() for c in row)
    wb.close()


def test_31_person_result_fields_and_complete_daily_rows(tmp_path: Path):
    wb = load_workbook(_export(tmp_path), data_only=True)
    detail = wb["工资核算明细"]
    assert detail.max_row == 34
    assert detail.cell(4, 6).value == 0
    assert detail.cell(4, 17).value == 90
    assert detail.cell(4, 18).value == 0
    assert wb["考勤汇总"].max_row == 3 + 31
    assert wb["人工复核事项"].cell(4, 3).value == "无待处理事项"
    wb.close()


@pytest.mark.parametrize("count", [33, 42])
def test_attendance_summary_is_one_row_per_person(tmp_path: Path, count: int):
    wb = load_workbook(_export(tmp_path, _fixture(count)), data_only=True)
    sheet = wb["考勤汇总"]
    assert sheet.max_row == 3 + count
    assert sheet.max_column == 14 + 30
    assert len({sheet.cell(row, 2).value for row in range(4, sheet.max_row + 1)}) == count
    wb.close()


@pytest.mark.parametrize(("month", "days"), [("2026-01", 31), ("2026-02", 28), ("2024-02", 29)])
def test_attendance_date_columns_follow_month_length(tmp_path: Path, month: str, days: int):
    wb = load_workbook(_export(tmp_path, _fixture(1, month)), data_only=True)
    sheet = wb["考勤汇总"]
    assert sheet.max_column == 14 + days
    assert [sheet.cell(3, 14 + day).value for day in range(1, days + 1)] == [f"{day}日" for day in range(1, days + 1)]
    assert sheet.max_row == 4
    wb.close()


def test_attendance_day_cell_retains_status_value_and_raw_mark(tmp_path: Path):
    data = _fixture(1)
    data["attendance_snapshot"]["records"][0]["items"][0].update(value="0.5", raw_value="8*", status=1)
    data["_source_people"][0]["daily"][0].update(value="0.5", raw_mark="8*")
    wb = load_workbook(_export(tmp_path, data), data_only=True)
    assert wb["考勤汇总"]["O4"].value == "出勤\n值：0.5\n原始：8*"
    wb.close()


def test_duplicate_names_remain_separate_by_record_id(tmp_path: Path):
    data = _fixture(2)
    duplicate_name = "同名合成人员"
    data["attendance_snapshot"]["records"][0]["name_snapshot"] = duplicate_name
    data["attendance_snapshot"]["records"][1]["name_snapshot"] = duplicate_name
    data["calculation_result"]["items"][0]["name"] = duplicate_name
    data["calculation_result"]["items"][1]["name"] = duplicate_name
    data["_source_people"][0]["name"] = duplicate_name
    data["_source_people"][1]["name"] = duplicate_name
    wb = load_workbook(_export(tmp_path, data), data_only=True)
    sheet = wb["考勤汇总"]
    assert sheet.cell(4, 1).value == sheet.cell(5, 1).value
    assert sheet.cell(4, 2).value != sheet.cell(5, 2).value
    wb.close()


def test_historical_template_style_and_money_formats_are_kept(tmp_path: Path):
    wb = load_workbook(_export(tmp_path), data_only=False)
    detail = wb["工资核算明细"]
    assert detail.column_dimensions["F"].width == 18
    assert detail["A1"].font.name == "Arial" and detail["A1"].font.sz == 15
    assert detail["A1"].fill.fgColor.rgb.endswith("D9EAF7")
    assert detail["F4"].number_format == exporter.MONEY_FMT
    assert wb["考勤汇总"]["O4"].value == "出勤\n值：（空）\n原始：8"
    wb.close()


def test_zero_source_and_missing_source_are_distinct(tmp_path: Path):
    data = _fixture(1)
    comps = data["calculation_result"]["items"][0]["components"]
    comps[0]["source_amount"] = "0.00"
    comps[1].update(source_present=False, source_amount=None)
    wb_path = _export(tmp_path, data)
    wb = load_workbook(wb_path, data_only=True)
    row = wb["工资基准"]
    assert row.cell(4, 8).value == 0 and row.cell(4, 9).value == "是"
    assert row.cell(4, 10).value is None and row.cell(4, 11).value == "否"
    assert row.cell(4, 12).value == 11.25 and row.cell(4, 13).value == "是"
    wb.close()


def test_false_null_zero_source_is_retained_without_redundant_review(tmp_path: Path):
    data = _fixture(1)
    data["calculation_result"]["items"][0]["components"][0].update(
        source_present=False, source_amount=None, calculated_amount="0.00")
    wb = load_workbook(_export(tmp_path, data), data_only=True)
    baseline = wb["工资基准"]
    assert baseline.cell(4, 9).value == "否" and baseline.cell(4, 8).value is None
    assert not any("来源缺失" in str(row[2]) or "来源与计算" in str(row[2])
                   for row in wb["人工复核事项"].iter_rows(min_row=4, values_only=True))
    wb.close()


def test_unknown_sources_aggregate_by_person_and_category(tmp_path: Path):
    data = _fixture(1)
    components = data["calculation_result"]["items"][0]["components"]
    components[0].update(source_present=None, source_amount=None)
    components[1].update(source_present=None, source_amount=None)
    _, gate_path, _ = _write_inputs(tmp_path, data)
    reviews = exporter._build_rows(data, gate_path)[-1]
    reviews = [row for row in reviews if row[2] == "工资分项来源状态未知"]
    assert len(reviews) == 1
    assert "基础工资" in reviews[0][3] and "绩效工资" in reviews[0][3]


def test_missing_source_with_nonzero_calculation_is_reported(tmp_path: Path):
    data = _fixture(1)
    data["calculation_result"]["items"][0]["components"][1].update(source_present=False, source_amount=None)
    wb = load_workbook(_export(tmp_path, data), data_only=True)
    reviews = [row for row in wb["人工复核事项"].iter_rows(min_row=4, values_only=True)
               if row[2] == "工资来源与计算金额不一致"]
    assert len(reviews) == 1 and "绩效工资" in reviews[0][3]
    wb.close()


def test_review_items_group_by_person_and_category_and_preserve_details(tmp_path: Path):
    data = _fixture(2)
    data["attendance_snapshot"]["records"][0]["reconciliation_status"] = 2
    data["attendance_snapshot"]["records"][0]["remark"] = "用户确认调整"
    data["calculation_result"]["items"][0]["components"][0].update(source_present=None, source_amount=None)
    data["calculation_result"]["items"][0]["components"][1].update(source_present=None, source_amount=None)
    data["calculation_result"]["issues"] = [
        {"payroll_person_id": 1, "code": "PAYROLL_NOTE", "message": "事项甲"},
        {"payroll_person_id": 1, "code": "PAYROLL_NOTE", "message": "事项乙"},
        {"payroll_person_id": 2, "code": "PAYROLL_NOTE", "message": "其他人员事项"},
    ]
    _, gate_path, _ = _write_inputs(tmp_path, data)
    rows = exporter._build_rows(data, gate_path)[-1]
    assert sum(1 for row in rows if row[0] == "合成人员01" and row[2] == "工资分项来源状态未知") == 1
    assert any(row[0] == "合成人员01" and row[2] == "考勤人工调整" for row in rows)
    payroll_note = [row for row in rows if row[2] == "PAYROLL_NOTE"]
    assert len(payroll_note) == 2
    assert any("事项甲" in row[3] and "事项乙" in row[3] for row in payroll_note)
    assert any(row[0] == "合成人员02" and "其他人员事项" in row[3] for row in payroll_note)


def test_same_name_review_issues_stay_separate_by_payroll_identity():
    attendance = [
        {"record_id": "record-a", "payroll_person_id": 101, "name_snapshot": "同名人员",
         "reconciliation_status": 2, "remark": "甲记录调整", "items": []},
        {"record_id": "record-b", "payroll_person_id": 202, "name_snapshot": "同名人员",
         "reconciliation_status": 2, "remark": "乙记录调整", "items": []},
    ]
    rows = exporter._review_items(attendance, {}, MONTH, [])
    adjustment_rows = [row for row in rows if row[3] == "考勤人工调整"]
    assert len(adjustment_rows) == 2
    assert adjustment_rows[0][0] != adjustment_rows[1][0]
    assert adjustment_rows[0][4] != adjustment_rows[1][4]


def test_missing_historical_wage_in_local_result_without_raw_dto_is_blocked(tmp_path: Path):
    data = _fixture(1)
    item = data["calculation_result"]["items"][0]
    item.update(his_salary_record_id=None, historical_source_month=None,
                historical_gross_amount=None, difference=None)
    data["calculation_result"]["summary"].update(historical_gross_total="0.00", difference_total="0.00")
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    output = tmp_path / "missing-history.xlsx"
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


def test_attendance_status_uses_payroll_adjusted_rest_label(tmp_path: Path):
    data = _fixture(1)
    data["attendance_snapshot"]["records"][0]["items"][0].update(status=6, value=None, raw_value="调")
    data["attendance_snapshot"]["records"][0].update(derived_actual_work_days=19, reconciliation_status=2)
    data["_source_people"][0]["daily"][0].update(status="comp_leave", raw_mark="调")
    data["_source_people"][0]["daily"][0]["user_confirmed_status"] = "comp_leave"
    wb = load_workbook(_export(tmp_path, data), data_only=True)
    assert wb["考勤汇总"].cell(4, 15).value == "调整休息\n值：（空）\n原始：调"
    wb.close()


@pytest.mark.parametrize("mutation", [
    lambda d: d["calculation_result"].update(status="PARTIAL"),
    lambda d: d["calculation_result"].update(status="FAILED"),
    lambda d: d["calculation_result"]["items"].pop(),
    lambda d: d["calculation_result"]["items"][1].update(payroll_person_id=d["calculation_result"]["items"][0]["payroll_person_id"]),
    lambda d: d["calculation_result"]["items"][0]["components"].pop(),
    lambda d: d["calculation_result"]["items"][0]["components"].append(copy.deepcopy(d["calculation_result"]["items"][0]["components"][0])),
    lambda d: d["calculation_result"].update(payroll_month="2026-05"),
    lambda d: d["calculation_result"].update(batch_id="another"),
    lambda d: d["latest_precheck"].update(valid=False),
    lambda d: d["calculation_result"]["items"][0]["components"][0].update(calculated_amount="NaN"),
    lambda d: d["calculation_result"]["items"][0]["components"][0].update(source_present=1),
    lambda d: d["calculation_result"]["items"][0]["components"][0].update(source_present=True, source_amount=None),
    lambda d: d["attendance_snapshot"]["records"][0]["items"][0].update(status=99),
])
def test_invalid_or_partial_inputs_fail_without_output(tmp_path: Path, mutation):
    data = _fixture(2)
    mutation(data)
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    output = tmp_path / "blocked.xlsx"
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


def test_gate_rejection_blocks_and_input_files_are_unchanged(tmp_path: Path):
    data = _fixture(1)
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    before = input_path.read_bytes()
    gate = json.loads(gate_path.read_text(encoding="utf-8"))
    gate["evidence"]["calculation_receipt"]["operation_id"] = "calculate:wrong"
    gate_path.write_text(json.dumps(gate), encoding="utf-8")
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, tmp_path / "blocked.xlsx")
    assert input_path.read_bytes() == before
    assert not (tmp_path / "blocked.xlsx").exists()


def test_non_calculated_attendance_person_is_preserved_and_flagged(tmp_path: Path):
    data = _fixture(1)
    extra = copy.deepcopy(data["attendance_snapshot"]["records"][0])
    extra["record_id"] = "record-unmatched"
    extra["payroll_person_id"] = None
    extra["roster_id"] = None
    extra["name_snapshot"] = "未匹配合成人员"
    data["attendance_snapshot"]["records"].append(extra)
    source_extra = copy.deepcopy(data["_source_people"][0])
    source_extra.update(record_id="record-unmatched", name="未匹配合成人员")
    source_extra["row_index"] = 2
    source_extra["printed_work_days_row_index"] = 2
    data["_source_people"].append(source_extra)
    data["_source_matches"].append({"record_id": "record-unmatched", "status": "unmatched",
                                    "payroll_person_id": None, "roster_id": None})
    wb = load_workbook(_export(tmp_path, data), data_only=True)
    attendance = wb["考勤汇总"]
    assert attendance.max_row == 3 + 2
    assert sum(1 for row in attendance.iter_rows(min_row=4, values_only=True) if row[0] == "未匹配合成人员") == 1
    assert any(row[0] == "未匹配合成人员" and row[2] == "未纳入核算"
               for row in wb["人工复核事项"].iter_rows(min_row=4, values_only=True))
    wb.close()


def _out_of_scope_unknown_fixture() -> dict:
    data = _fixture(42)
    data["calculation_result"]["items"] = data["calculation_result"]["items"][:31]
    result = data["calculation_result"]
    result.update(population_count=31, participant_count=31, calculated_count=31)
    result["summary"].update(population_count=31, calculated_count=31,
                             historical_gross_total="2790.00", calculated_gross_total="2790.00")
    data["latest_precheck"].update(participant_count=31, calculated_count=31)

    source = data["_source_people"][31]
    source.update(name="程瑶", printed_work_days=15)
    source["daily"] = [
        {"day": day, "status": "work" if day <= 15 else "departed",
         "raw_mark": "8" if day <= 15 else "离"}
        for day in range(1, 31)
    ]
    snapshot = data["attendance_snapshot"]["records"][31]
    snapshot.update(name_snapshot="程瑶", source_actual_work_days=15,
                    derived_actual_work_days=15, confirmed_actual_work_days=15)
    for item in snapshot["items"]:
        if item["day"] >= 16:
            item.update(status=None, raw_value="离")
    data["_source_matches"][31].update(name="程瑶")
    return data


def test_42_attendance_people_31_calculated_preserves_out_of_scope_unknown_source(tmp_path: Path):
    data = _out_of_scope_unknown_fixture()
    input_path, gate_path, attendance_path = _write_inputs(tmp_path, data)
    original_input = input_path.read_bytes()
    original_attendance = attendance_path.read_bytes()

    result = _call_export(input_path, gate_path, tmp_path / "42-attendance-31-calculated.xlsx")

    assert result["calculated_people"] == 31
    assert result["attendance_people_rows"] == 42
    assert input_path.read_bytes() == original_input
    assert attendance_path.read_bytes() == original_attendance
    wb = load_workbook(result["path"], data_only=True)
    assert wb.sheetnames == list(exporter.SHEETS)
    assert [wb[name].max_column for name in wb.sheetnames] == [19, 25, 44, 8, 4]
    detail = wb["工资核算明细"]
    assert detail.max_row == 34
    assert detail.cell(4, 17).value == 90
    attendance = wb["考勤汇总"]
    assert attendance.max_row == 45
    cheng = next(row for row in attendance.iter_rows(min_row=4, values_only=True) if row[0] == "程瑶")
    assert cheng[1] == "record-32"
    assert (cheng[6], cheng[7], cheng[8]) == (15, 15, 15)
    assert cheng[29] == "待核实（来源状态：离职后）\n值：（空）\n原始：离"
    submission = _read_json(tmp_path / "submission-request.json")
    assert all(record["payroll_person_id"] != 32 for record in submission["records"])
    review_rows = list(wb["人工复核事项"].iter_rows(min_row=4, values_only=True))
    assert any(row[0] == "程瑶" and row[2] == "核算范围外考勤状态待核实" and "16日" in row[3] for row in review_rows)
    assert any(row[0] == "程瑶" and row[2] == "未纳入核算" for row in review_rows)
    wb.close()


def test_hanyue_42_to_31_uses_confirmed_payroll_project_and_preserves_source_project(tmp_path: Path):
    data = _out_of_scope_unknown_fixture()
    source_project, payroll_project = "瀚阅府西苑项目", "瀚阅府"
    for person in data["_source_people"]:
        person["project"] = source_project
    for record in data["attendance_snapshot"]["records"]:
        record["source_project_snapshot"] = source_project
        record["project_snapshot"] = payroll_project
    for item in data["calculation_result"]["items"]:
        item["project_name"] = payroll_project

    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    result = _call_export(input_path, gate_path, tmp_path / "hanyue-confirmed-project.xlsx")

    assert result["calculated_people"] == 31
    assert result["attendance_people_rows"] == 42
    submission = _read_json(tmp_path / "submission-request.json")
    assert len(submission["records"]) == 31
    assert {row["source_project_snapshot"] for row in submission["records"]} == {source_project}
    assert {row["project_snapshot"] for row in submission["records"]} == {payroll_project}
    wb = load_workbook(result["path"], data_only=True)
    assert wb["工资核算明细"].max_row == 34
    assert wb["考勤汇总"].max_row == 45
    assert wb["工资核算明细"].cell(4, 2).value == payroll_project
    assert wb["考勤汇总"].cell(4, 3).value == payroll_project
    wb.close()


def test_unmapped_attendance_status_still_blocks_calculated_person(tmp_path: Path):
    data = _fixture(1)
    source_day = data["_source_people"][0]["daily"][15]
    source_day.update(status="departed", raw_mark="离")
    data["attendance_snapshot"]["records"][0]["items"][15].update(status=7, raw_value="离")
    data["attendance_snapshot"]["records"][0].update(derived_actual_work_days=19)
    input_path, gate_path, _ = _write_inputs(tmp_path, data)

    with pytest.raises(exporter.ExportError, match="ATTENDANCE_STATUS_HAS_NO_PAYROLL_EQUIVALENT"):
        _call_export(input_path, gate_path, tmp_path / "included-unknown.xlsx")
    assert not (tmp_path / "included-unknown.xlsx").exists()


def test_out_of_scope_unknown_status_does_not_skip_snapshot_source_checks(tmp_path: Path):
    data = _out_of_scope_unknown_fixture()
    data["attendance_snapshot"]["records"][31]["items"][15]["raw_value"] = "错误标记"
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    output = tmp_path / "tampered-out-of-scope.xlsx"

    with pytest.raises(exporter.ExportError, match="ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


@pytest.mark.parametrize("status", [1, 5, 7])
def test_out_of_scope_unmapped_status_requires_null_snapshot_placeholder(tmp_path: Path, status: int):
    data = _out_of_scope_unknown_fixture()
    data["attendance_snapshot"]["records"][31]["items"][15]["status"] = status
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    output = tmp_path / f"tampered-status-{status}.xlsx"

    with pytest.raises(exporter.ExportError, match="ATTENDANCE_SNAPSHOT_SOURCE_MISMATCH"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


def test_out_of_scope_invalid_source_status_is_not_treated_as_displayable_unknown(tmp_path: Path):
    data = _out_of_scope_unknown_fixture()
    data["_source_people"][31]["daily"][15]["status"] = "invented_status"
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    output = tmp_path / "invalid-out-of-scope-status.xlsx"

    with pytest.raises(exporter.ExportError, match="ATTENDANCE_SOURCE_STATUS_INVALID"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


def test_repeated_export_has_same_business_cells_and_never_overwrites_input(tmp_path: Path):
    data = _fixture(2)
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    before = input_path.read_bytes()
    one = _call_export(input_path, gate_path, tmp_path / "one.xlsx")
    two = _call_export(input_path, gate_path, tmp_path / "two.xlsx")
    assert input_path.read_bytes() == before
    w1, w2 = load_workbook(one["path"], data_only=True), load_workbook(two["path"], data_only=True)
    for name in w1.sheetnames:
        assert list(w1[name].values) == list(w2[name].values)
    w1.close()
    w2.close()


def test_existing_output_is_never_overwritten(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    output = tmp_path / "already.xlsx"
    output.write_bytes(b"keep")
    with pytest.raises(exporter.ExportError, match="OUTPUT_PATH_INVALID_OR_EXISTS"):
        _call_export(input_path, gate_path, output)
    assert output.read_bytes() == b"keep"


def test_missing_or_corrupt_calculation_input_is_blocked(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    input_path.write_text("{broken", encoding="utf-8")
    with pytest.raises(exporter.ExportError, match="INPUT_UNREADABLE"):
        _call_export(input_path, gate_path, tmp_path / "bad-json.xlsx")
    input_path.unlink()
    with pytest.raises(exporter.ExportError, match="INPUT_UNREADABLE"):
        _call_export(input_path, gate_path, tmp_path / "missing.xlsx")
    assert not (tmp_path / "bad-json.xlsx").exists()
    assert not (tmp_path / "missing.xlsx").exists()


def test_original_calculation_dto_blocks_self_consistent_tampering(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    data = _read_json(input_path)
    result = data["calculation_result"]
    item = result["items"][0]
    item["components"][0]["calculated_amount"] = "1.00"
    item["calculated_gross_amount"] = "91.00"
    item["difference"] = "1.00"
    result["summary"]["calculated_gross_total"] = "91.00"
    result["summary"]["difference_total"] = "1.00"
    _write_json(input_path, data)
    output = tmp_path / "tampered-internal.xlsx"
    with pytest.raises(exporter.ExportError, match="CALCULATION_RESULT_SOURCE_MISMATCH"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


@pytest.mark.parametrize("mutate", [
    lambda result: (result["items"][0].update(historical_gross_amount="91.00", difference="-1.00"),
                    result["summary"].update(historical_gross_total="91.00", difference_total="-1.00")),
    lambda result: result["items"][0]["components"][0].update(source_present=False, source_amount=None),
    lambda result: result["items"].pop(),
    lambda result: result["items"][0]["components"][0].update(source_amount="999.00"),
])
def test_calculation_field_changes_are_blocked_against_raw_response(tmp_path: Path, mutate):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    data = _read_json(input_path)
    mutate(data["calculation_result"])
    _write_json(input_path, data)
    output = tmp_path / "changed-calculation.xlsx"
    with pytest.raises(exporter.ExportError, match="CALCULATION_RESULT_SOURCE_MISMATCH"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


@pytest.mark.parametrize("mutation", [
    lambda rec: rec["items"][0].update(status=5),
    lambda rec: rec["items"][0].update(raw_value="伪造"),
    lambda rec: rec.update(confirmed_actual_work_days="19"),
    lambda rec: rec.update(payroll_person_id=999),
])
def test_attendance_snapshot_must_match_current_confirmed_source(tmp_path: Path, mutation):
    input_path, gate_path, attendance_path = _write_inputs(tmp_path, _fixture(1))
    data = _read_json(input_path)
    mutation(data["attendance_snapshot"]["records"][0])
    _write_json(input_path, data)
    output = tmp_path / "changed-attendance.xlsx"
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, output)
    assert not output.exists()
    assert attendance_path.exists()


def test_snapshot_cannot_drop_an_attendance_record(tmp_path: Path):
    data = _fixture(2)
    data["attendance_snapshot"]["records"].pop()
    input_path, gate_path, _ = _write_inputs(tmp_path, data)
    output = tmp_path / "dropped-attendance.xlsx"
    with pytest.raises(exporter.ExportError, match="ATTENDANCE_SNAPSHOT_SOURCE_SCOPE_MISMATCH"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


def test_current_attendance_sha_must_match_review_and_identity_mapping(tmp_path: Path):
    input_path, gate_path, attendance_path = _write_inputs(tmp_path, _fixture(1))
    attendance = _read_json(attendance_path)
    attendance["people"][0]["name"] = "修改后姓名"
    _write_json(attendance_path, attendance)
    output = tmp_path / "stale-attendance.xlsx"
    with pytest.raises(exporter.ExportError, match="ATTENDANCE_REVIEW_SOURCE_MISMATCH"):
        _call_export(input_path, gate_path, output)
    assert not output.exists()


def test_canonical_json_comparison_ignores_whitespace_and_key_order(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    raw_path = tmp_path / "calculation-response.json"
    _write_json(raw_path, _read_json(raw_path), indent=4, sort_keys=True)
    result = _call_export(input_path, gate_path, tmp_path / "canonical.xlsx")
    assert result["ok"] is True


def test_payroll_mcp_call_tool_result_is_normalized_without_changing_raw_file(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    raw_path = tmp_path / "calculation-response.json"
    dto = _read_json(raw_path)
    dto["issues"] = None  # Go nil slice serializes as JSON null; equivalent to no issues.
    text = json.dumps(dto, ensure_ascii=False, separators=(",", ":"))
    wrapped = {
        "content": [{"type": "text", "text": text}],
        "structuredContent": dto,
    }
    raw_path.write_text(json.dumps(wrapped, ensure_ascii=False), encoding="utf-8")
    before = raw_path.read_bytes()

    result = _call_export(input_path, gate_path, tmp_path / "mcp-result.xlsx")

    assert result["ok"] is True
    assert raw_path.read_bytes() == before


def test_payroll_mcp_text_and_structured_payload_must_match(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    raw_path = tmp_path / "calculation-response.json"
    dto = _read_json(raw_path)
    text_dto = copy.deepcopy(dto)
    text_dto["items"][0]["calculated_gross_amount"] = "999.00"
    raw_path.write_text(json.dumps({
        "content": [{"type": "text", "text": json.dumps(text_dto, ensure_ascii=False)}],
        "structuredContent": dto,
    }, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(exporter.ExportError, match="RAW_CALCULATION_MCP_CONTENT_MISMATCH"):
        _call_export(input_path, gate_path, tmp_path / "mcp-conflict.xlsx")
    assert not (tmp_path / "mcp-conflict.xlsx").exists()


@pytest.mark.parametrize("wrapper", [
    {"structuredContent": {"payroll_month": "2026-06"}, "content": []},
    {"structuredContent": {}, "content": [{"type": "text", "text": "{}"}], "isError": True},
    {"structuredContent": {}, "content": [{"type": "text", "text": "{}"}], "unexpected": True},
])
def test_malformed_or_unrecognized_mcp_wrappers_fail_closed(tmp_path: Path, wrapper: dict):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    _write_json(tmp_path / "calculation-response.json", wrapper)
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, tmp_path / "bad-wrapper.xlsx")
    assert not (tmp_path / "bad-wrapper.xlsx").exists()


def test_optional_attendance_id_and_null_source_amount_follow_go_dto_contract(tmp_path: Path):
    data = _fixture(1)
    component = data["calculation_result"]["items"][0]["components"][0]
    component.update(source_present=False, source_amount=None)
    data["calculation_result"]["items"][0].pop("attendance_id", None)
    path = _export(tmp_path, data)
    assert path.exists()


def test_null_issues_normalize_but_null_items_and_missing_required_keys_fail(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    raw_path = tmp_path / "calculation-response.json"
    dto = _read_json(raw_path)
    dto["issues"] = None
    raw_path.write_text(json.dumps(dto, ensure_ascii=False), encoding="utf-8")
    data = _read_json(input_path)
    data["calculation_result"]["issues"] = None
    _write_json(input_path, data)
    assert _call_export(input_path, gate_path, tmp_path / "null-issues.xlsx")["ok"] is True

    dto["items"] = None
    _write_json(raw_path, dto)
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, tmp_path / "null-items.xlsx")
    assert not (tmp_path / "null-items.xlsx").exists()

    dto["items"] = []
    _write_json(raw_path, dto)
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, tmp_path / "empty-items.xlsx")
    assert not (tmp_path / "empty-items.xlsx").exists()


def test_raw_calculation_response_missing_corrupt_or_wrapped_is_blocked(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    raw_path = tmp_path / "calculation-response.json"
    original = raw_path.read_bytes()
    raw_path.unlink()
    with pytest.raises(exporter.ExportError, match="RAW_CALCULATION_RESPONSE_UNREADABLE"):
        _call_export(input_path, gate_path, tmp_path / "missing-response.xlsx")
    raw_path.write_text("{broken", encoding="utf-8")
    with pytest.raises(exporter.ExportError):
        _call_export(input_path, gate_path, tmp_path / "corrupt-response.xlsx")
    _write_json(raw_path, {"result": json.loads(original)})
    with pytest.raises(exporter.ExportError, match="RAW_CALCULATION_RESPONSE_STRUCTURE_INVALID"):
        _call_export(input_path, gate_path, tmp_path / "wrapped-response.xlsx")


def test_unknown_calculation_dto_field_is_not_silently_discarded(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    raw_path = tmp_path / "calculation-response.json"
    raw = _read_json(raw_path)
    raw["unexpected_field"] = "not silently ignored"
    _write_json(raw_path, raw)
    with pytest.raises(exporter.ExportError, match="RAW_CALCULATION_RESPONSE_STRUCTURE_INVALID"):
        _call_export(input_path, gate_path, tmp_path / "unknown-field.xlsx")


def test_submission_request_must_match_snapshot_and_receipt_scope(tmp_path: Path):
    input_path, gate_path, _ = _write_inputs(tmp_path, _fixture(1))
    request_path = tmp_path / "submission-request.json"
    request = _read_json(request_path)
    request["records"][0]["items"][0]["raw_value"] = "伪造"
    _write_json(request_path, request)
    with pytest.raises(exporter.ExportError, match="SUBMISSION_REQUEST_SNAPSHOT_MISMATCH"):
        _call_export(input_path, gate_path, tmp_path / "request-mismatch.xlsx")
    assert not (tmp_path / "request-mismatch.xlsx").exists()


def test_successful_export_never_modifies_source_files(tmp_path: Path):
    input_path, gate_path, attendance_path = _write_inputs(tmp_path, _fixture(1))
    paths = [input_path, gate_path, attendance_path, tmp_path / "calculation-response.json",
             tmp_path / "person-matches.json", tmp_path / "submission-request.json"]
    before = {path: path.read_bytes() for path in paths}
    _call_export(input_path, gate_path, tmp_path / "read-only.xlsx")
    assert {path: path.read_bytes() for path in paths} == before


def test_historical_success_workbook_matches_payroll_result_without_modifying_it():
    historical = Path("/Users/ccy/.leros/projects/1/prj_Z31vkWS1d1/repo")
    input_path = historical / ".payroll-workbook-input.json"
    workbook_path = historical / "2026年6月工资核算表.xlsx"
    if not input_path.is_file() or not workbook_path.is_file():
        pytest.skip("historical read-only sample is unavailable")
    before = (input_path.stat().st_mtime_ns, workbook_path.stat().st_mtime_ns)
    data = json.loads(input_path.read_text(encoding="utf-8"))
    book = load_workbook(workbook_path, data_only=True, read_only=True)
    ws = book["工资核算明细"]
    source = {row["name"]: row for row in data["calculation_result"]["items"]}
    code_order = [code for code, _ in COMPONENTS]
    checked = 0
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row[0]:
            continue
        item = source[row[0]]
        values = {component["code"]: Decimal(component["calculated_amount"]) for component in item["components"]}
        assert [Decimal(str(value)) for value in row[5:14]] == [values[code] for code in code_order]
        assert Decimal(str(row[15])) == Decimal(item["historical_gross_amount"])
        assert Decimal(str(row[16])) == Decimal(item["calculated_gross_amount"])
        assert Decimal(str(row[17])) == Decimal(item["difference"])
        checked += 1
    book.close()
    assert checked == 31
    assert before == (input_path.stat().st_mtime_ns, workbook_path.stat().st_mtime_ns)
